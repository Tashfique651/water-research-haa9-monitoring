"""Species-level UCMR4 composition addendum using frozen Peterson FOIA data.

Official EPA HAA5, HAA6Br, and HAA9 group totals remain the primary group
measurements. Species records are used for attribution and censor-aware
sensitivity only. The script writes analysis tables, source data, a machine
summary, and a publication-ready supplementary figure.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.lines import Line2D
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PET = ROOT / "data" / "raw" / "additional" / "peterson_haa"
FOIA = PET / "python" / "data" / "UCMR4 FOIA"
UCMR4 = ROOT / "data" / "raw" / "ucmr4-occurrence-data.zip"
OUT = ROOT / "results" / "additional"
FIG = ROOT / "figures" / "supplementary"

COMMIT = "e0dd19e8f392ac08aa3a00e6c09622014d074e7c"
SEED = 20260913
BOOTSTRAPS = 10_000

SPECIES = {
    "monobromoacetic acid": "MBAA",
    "monochloroacetic acid": "MCAA",
    "dibromoacetic acid": "DBAA",
    "dichloroacetic acid": "DCAA",
    "bromochloroacetic acid": "BCAA",
    "tribromoacetic acid": "TBAA",
    "trichloroacetic acid": "TCAA",
    "chlorodibromoacetic acid": "CDBAA",
    "bromodichloroacetic acid": "BDCAA",
}
FULL_NAME = {v: k for k, v in SPECIES.items()}
ORDER = ["MCAA", "DCAA", "TCAA", "MBAA", "DBAA", "BCAA", "BDCAA", "CDBAA", "TBAA"]
HAA5 = {"MBAA", "MCAA", "DBAA", "DCAA", "TCAA"}
EXTRA4 = {"BCAA", "BDCAA", "CDBAA", "TBAA"}
HAA6BR = {"MBAA", "DBAA", "BCAA", "BDCAA", "CDBAA", "TBAA"}
NONBR3 = {"MCAA", "DCAA", "TCAA"}
GROUPS = {"Extra4": EXTRA4, "HAA6Br": HAA6BR}

COLORS = {
    "MBAA": "#56B4E9",
    "DBAA": "#009E73",
    "BCAA": "#E69F00",
    "BDCAA": "#D55E00",
    "CDBAA": "#CC79A7",
    "TBAA": "#6F4C9B",
}


def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def iso_date(value) -> str:
    s = clean(value)
    if not s:
        return ""
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S.%f"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return s


def norm_event(value) -> str:
    return re.sub(r"^SEH", "SE", clean(value).upper())


def event_key(pws, facility, point, sample_id, date, event) -> tuple[str, ...]:
    return (
        clean(pws), clean(facility), clean(point), clean(sample_id),
        iso_date(date), norm_event(event),
    )


def as_float(value) -> float:
    s = clean(value)
    return float(s) if s else math.nan


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def save_table(df: pd.DataFrame, stem: str) -> tuple[Path, Path]:
    csv_path = OUT / f"{stem}.csv"
    parquet_path = OUT / f"{stem}.parquet"
    df.to_csv(csv_path, index=False, float_format="%.10g")
    df.to_parquet(parquet_path, index=False, compression="zstd")
    return csv_path, parquet_path


def load_species_events() -> tuple[dict, dict]:
    files = sorted(FOIA.glob("*.xlsx"))
    events: dict[tuple[str, ...], dict] = {}
    rows_by_file = Counter()
    signs = Counter()
    methods = Counter()
    duplicates = 0
    headers = None
    for path in files:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        raw_header = [clean(v) for v in next(rows)]
        headers = raw_header
        for row in rows:
            if not row or all(v is None for v in row):
                continue
            rows_by_file[path.name] += 1
            (
                pwsid, pws_name, size, facility_id, facility_name,
                facility_water_type, sample_point_id, sample_point_name,
                sample_point_type, collection_date, sample_id, contaminant,
                mrl, method_id, sign, value, sample_event_code,
                monitoring_requirement, region, state,
            ) = row
            abbr = SPECIES.get(clean(contaminant).lower())
            if abbr is None:
                raise ValueError(f"Unexpected species {contaminant!r} in {path}")
            key = event_key(
                pwsid, facility_id, sample_point_id, sample_id,
                collection_date, sample_event_code,
            )
            if key not in events:
                events[key] = {
                    "PWSID": clean(pwsid),
                    "PWSName": clean(pws_name),
                    "Size": clean(size),
                    "FacilityID": clean(facility_id),
                    "FacilityName": clean(facility_name),
                    "FacilityWaterType": clean(facility_water_type),
                    "SamplePointID": clean(sample_point_id),
                    "SamplePointName": clean(sample_point_name),
                    "SamplePointType": clean(sample_point_type),
                    "CollectionDate": iso_date(collection_date),
                    "SampleID": clean(sample_id),
                    "SampleEventCode": norm_event(sample_event_code),
                    "MonitoringRequirement": clean(monitoring_requirement),
                    "Region": clean(region),
                    "State": clean(state),
                    "species": {},
                }
            if abbr in events[key]["species"]:
                duplicates += 1
            detected = clean(sign) == "="
            numeric = as_float(value)
            if detected and not np.isfinite(numeric):
                raise ValueError(f"Detected {abbr} lacks a value at {key}")
            mrl_value = as_float(mrl)
            events[key]["species"][abbr] = {
                "sign": clean(sign),
                "value": numeric if detected else math.nan,
                "mrl": mrl_value,
                "method": clean(method_id),
            }
            signs[clean(sign)] += 1
            methods[clean(method_id)] += 1
        wb.close()
    audit = {
        "files": [str(p) for p in files],
        "headers": headers,
        "rows_by_file": dict(rows_by_file),
        "rows": sum(rows_by_file.values()),
        "events": len(events),
        "systems": len({k[0] for k in events}),
        "panel_size_counts": dict(Counter(len(v["species"]) for v in events.values())),
        "duplicates": duplicates,
        "sign_counts": dict(signs),
        "method_counts": dict(methods),
    }
    return events, audit


def load_official_groups() -> tuple[dict, dict]:
    groups: dict[tuple[str, ...], dict[str, float]] = defaultdict(dict)
    rows = Counter()
    signs = Counter()
    with zipfile.ZipFile(UCMR4) as z, z.open("UCMR4_All.txt") as b:
        text_rows = (line.decode("utf-8-sig", "replace") for line in b)
        reader = csv.DictReader(text_rows, delimiter="\t")
        for row in reader:
            group = row["Contaminant"]
            if group not in {"HAA5", "HAA6Br", "HAA9"}:
                continue
            key = event_key(
                row["PWSID"], row["FacilityID"], row["SamplePointID"],
                row["SampleID"], row["CollectionDate"], row["SampleEventCode"],
            )
            groups[key][group] = as_float(row["AnalyticalResultValue"])
            rows[group] += 1
            signs[row["AnalyticalResultsSign"]] += 1
    audit = {
        "rows_by_group": dict(rows),
        "events_union": len(groups),
        "systems_union": len({k[0] for k in groups}),
        "sign_counts": dict(signs),
    }
    return groups, audit


def build_event_frame(events: dict, groups: dict) -> tuple[pd.DataFrame, dict]:
    rows = []
    incomplete = Counter()
    missing_official = Counter()
    for key, event in events.items():
        panel = event["species"]
        if set(panel) != set(ORDER):
            incomplete[len(panel)] += 1
            continue
        official = groups.get(key, {})
        for group in ("HAA5", "HAA6Br", "HAA9"):
            if group not in official:
                missing_official[group] += 1
        if set(official) != {"HAA5", "HAA6Br", "HAA9"}:
            continue
        out = {k: v for k, v in event.items() if k != "species"}
        for abbr in ORDER:
            x = panel[abbr]
            detect = x["sign"] == "="
            lower = x["value"] if detect else 0.0
            upper = x["value"] if detect else x["mrl"]
            midpoint = x["value"] if detect else x["mrl"] / 2.0
            out[f"{abbr}_detected"] = detect
            out[f"{abbr}_mrl_ug_L"] = x["mrl"]
            out[f"{abbr}_lower_ug_L"] = lower
            out[f"{abbr}_mid_ug_L"] = midpoint
            out[f"{abbr}_upper_ug_L"] = upper
        out["HAA5_official_ug_L"] = official["HAA5"]
        out["HAA6Br_official_ug_L"] = official["HAA6Br"]
        out["HAA9_official_ug_L"] = official["HAA9"]
        out["Extra4_official_ug_L"] = official["HAA9"] - official["HAA5"]
        for suffix in ("lower", "mid", "upper"):
            out[f"HAA5_species_{suffix}_ug_L"] = sum(out[f"{s}_{suffix}_ug_L"] for s in HAA5)
            out[f"Extra4_species_{suffix}_ug_L"] = sum(out[f"{s}_{suffix}_ug_L"] for s in EXTRA4)
            out[f"HAA6Br_species_{suffix}_ug_L"] = sum(out[f"{s}_{suffix}_ug_L"] for s in HAA6BR)
            out[f"HAA9_species_{suffix}_ug_L"] = sum(out[f"{s}_{suffix}_ug_L"] for s in ORDER)
        h9 = out["HAA9_official_ug_L"]
        out["Extra4_fraction_official"] = out["Extra4_official_ug_L"] / h9 if h9 > 0 else math.nan
        out["HAA6Br_fraction_official"] = out["HAA6Br_official_ug_L"] / h9 if h9 > 0 else math.nan
        for target in ("Extra4", "HAA6Br"):
            for suffix in ("lower", "mid", "upper"):
                den = out[f"HAA9_species_{suffix}_ug_L"]
                out[f"{target}_fraction_{suffix}"] = out[f"{target}_species_{suffix}_ug_L"] / den if den > 0 else math.nan
        extra_low = out["Extra4_species_lower_ug_L"]
        extra_up = out["Extra4_species_upper_ug_L"]
        haa5_low = out["HAA5_species_lower_ug_L"]
        haa5_up = out["HAA5_species_upper_ug_L"]
        out["Extra4_fraction_interval_lower"] = extra_low / (extra_low + haa5_up) if extra_low + haa5_up > 0 else math.nan
        out["Extra4_fraction_interval_upper"] = extra_up / (extra_up + haa5_low) if extra_up + haa5_low > 0 else math.nan
        br_low = out["HAA6Br_species_lower_ug_L"]
        br_up = out["HAA6Br_species_upper_ug_L"]
        nonbr_low = sum(out[f"{s}_lower_ug_L"] for s in NONBR3)
        nonbr_up = sum(out[f"{s}_upper_ug_L"] for s in NONBR3)
        out["HAA6Br_fraction_interval_lower"] = br_low / (br_low + nonbr_up) if br_low + nonbr_up > 0 else math.nan
        out["HAA6Br_fraction_interval_upper"] = br_up / (br_up + nonbr_low) if br_up + nonbr_low > 0 else math.nan
        rows.append(out)
    df = pd.DataFrame(rows).sort_values(
        ["PWSID", "CollectionDate", "FacilityID", "SamplePointID", "SampleID"]
    ).reset_index(drop=True)
    audit = {
        "complete_species_events": int(sum(1 for v in events.values() if set(v["species"]) == set(ORDER))),
        "analysis_events": len(df),
        "analysis_systems": int(df["PWSID"].nunique()),
        "incomplete_panel_counts": dict(incomplete),
        "complete_events_missing_official_group": dict(missing_official),
        "date_min": df["CollectionDate"].min(),
        "date_max": df["CollectionDate"].max(),
    }
    return df, audit


def ratio_columns(frame: pd.DataFrame) -> pd.DataFrame:
    out = frame.copy()
    h9 = out["HAA9_official_ug_L"]
    out["Extra4_fraction_official"] = np.where(h9 > 0, out["Extra4_official_ug_L"] / h9, np.nan)
    out["HAA6Br_fraction_official"] = np.where(h9 > 0, out["HAA6Br_official_ug_L"] / h9, np.nan)
    for target in ("Extra4", "HAA6Br"):
        for suffix in ("lower", "mid", "upper"):
            den = out[f"HAA9_species_{suffix}_ug_L"]
            out[f"{target}_fraction_{suffix}"] = np.where(
                den > 0, out[f"{target}_species_{suffix}_ug_L"] / den, np.nan
            )
    extra_low = out["Extra4_species_lower_ug_L"]
    extra_up = out["Extra4_species_upper_ug_L"]
    haa5_low = out["HAA5_species_lower_ug_L"]
    haa5_up = out["HAA5_species_upper_ug_L"]
    out["Extra4_fraction_interval_lower"] = np.where(
        extra_low + haa5_up > 0, extra_low / (extra_low + haa5_up), np.nan
    )
    out["Extra4_fraction_interval_upper"] = np.where(
        extra_up + haa5_low > 0, extra_up / (extra_up + haa5_low), np.nan
    )
    br_low = out["HAA6Br_species_lower_ug_L"]
    br_up = out["HAA6Br_species_upper_ug_L"]
    nonbr_low = sum(out[f"{s}_lower_ug_L"] for s in NONBR3)
    nonbr_up = sum(out[f"{s}_upper_ug_L"] for s in NONBR3)
    out["HAA6Br_fraction_interval_lower"] = np.where(
        br_low + nonbr_up > 0, br_low / (br_low + nonbr_up), np.nan
    )
    out["HAA6Br_fraction_interval_upper"] = np.where(
        br_up + nonbr_low > 0, br_up / (br_up + nonbr_low), np.nan
    )
    return out


def build_system_frame(events: pd.DataFrame) -> pd.DataFrame:
    numeric = [c for c in events.columns if c.endswith("_ug_L")]
    detect_cols = [c for c in events.columns if c.endswith("_detected")]
    means = events.groupby("PWSID", sort=True)[numeric].mean()
    detect_any = events.groupby("PWSID", sort=True)[detect_cols].max().astype(bool)
    counts = events.groupby("PWSID", sort=True).size().rename("event_count")
    system = pd.concat([counts, means, detect_any], axis=1).reset_index()
    system = ratio_columns(system)
    ratio_names = [
        f"{target}_fraction_{suffix}"
        for target in ("Extra4", "HAA6Br")
        for suffix in ("official", "lower", "mid", "upper", "interval_lower", "interval_upper")
    ]
    mean_event_ratios = events.groupby("PWSID", sort=True)[ratio_names].mean()
    for col in ratio_names:
        system[f"{col}_ratio_of_means"] = system[col]
        system[col] = system["PWSID"].map(mean_event_ratios[col])
    return system


def species_coverage(events: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for abbr in ORDER:
        detected = events[f"{abbr}_detected"]
        values = events.loc[detected, f"{abbr}_lower_ug_L"]
        rows.append({
            "species": FULL_NAME[abbr],
            "abbreviation": abbr,
            "mrl_ug_L": float(events[f"{abbr}_mrl_ug_L"].drop_duplicates().iloc[0]),
            "events": len(events),
            "systems": int(events["PWSID"].nunique()),
            "detected_events": int(detected.sum()),
            "detected_event_fraction": float(detected.mean()),
            "censored_events": int((~detected).sum()),
            "systems_with_detection": int(events.loc[detected, "PWSID"].nunique()),
            "detected_mean_ug_L": float(values.mean()),
            "detected_median_ug_L": float(values.median()),
            "all_event_mean_nd0_ug_L": float(events[f"{abbr}_lower_ug_L"].mean()),
        })
    return pd.DataFrame(rows)


def bootstrap_ratios(cluster: pd.DataFrame, numerators: list[str], denominator: str, seed: int) -> dict[str, tuple[float, float]]:
    arr_num = cluster[numerators].to_numpy(dtype=float)
    arr_den = cluster[denominator].to_numpy(dtype=float)
    n = len(cluster)
    rng = np.random.default_rng(seed)
    draws = np.empty((BOOTSTRAPS, len(numerators)), dtype=float)
    batch = 100
    for start in range(0, BOOTSTRAPS, batch):
        stop = min(start + batch, BOOTSTRAPS)
        idx = rng.integers(0, n, size=(stop - start, n))
        den = arr_den[idx].sum(axis=1)
        for j in range(len(numerators)):
            draws[start:stop, j] = arr_num[idx, j].sum(axis=1) / den
    return {
        name: (float(np.quantile(draws[:, j], 0.025)), float(np.quantile(draws[:, j], 0.975)))
        for j, name in enumerate(numerators)
    }


def contribution_table(events: pd.DataFrame, systems: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for target_index, (target, members) in enumerate(GROUPS.items()):
        official = f"{target}_official_ug_L"
        for scope_index, (scope, frame) in enumerate((
            ("event_mass_weighted", events),
            ("equal_system_mass_weighted", systems),
        )):
            numerators = [f"{s}_lower_ug_L" for s in sorted(members)]
            if scope == "event_mass_weighted":
                cluster = events.groupby("PWSID", sort=True)[numerators + [official]].sum()
            else:
                cluster = systems.set_index("PWSID")[numerators + [official]]
            cis = bootstrap_ratios(
                cluster, numerators, official,
                SEED + 100 * target_index + 10 * scope_index,
            )
            totals = {
                suffix: {s: float(frame[f"{s}_{suffix}_ug_L"].sum()) for s in members}
                for suffix in ("lower", "mid", "upper")
            }
            official_total = float(frame[official].sum())
            for s in sorted(members):
                low = totals["lower"][s]
                mid = totals["mid"][s]
                up = totals["upper"][s]
                other_low = sum(totals["lower"][x] for x in members if x != s)
                other_up = sum(totals["upper"][x] for x in members if x != s)
                lower_bound = low / (low + other_up) if low + other_up > 0 else math.nan
                upper_bound = up / (up + other_low) if up + other_low > 0 else math.nan
                reconstructed_lower_total = sum(totals["lower"].values())
                reconstructed_mid_total = sum(totals["mid"].values())
                rows.append({
                    "target_group": target,
                    "species": FULL_NAME[s],
                    "abbreviation": s,
                    "weighting": scope,
                    "events": len(events),
                    "systems": len(systems),
                    "primary_detect_only_share_official_denominator": low / official_total,
                    "primary_cluster_bootstrap_ci_low": cis[f"{s}_lower_ug_L"][0],
                    "primary_cluster_bootstrap_ci_high": cis[f"{s}_lower_ug_L"][1],
                    "reconstructed_nd0_share": low / reconstructed_lower_total,
                    "midpoint_substitution_share": mid / reconstructed_mid_total,
                    "censor_interval_share_lower": lower_bound,
                    "censor_interval_share_upper": upper_bound,
                    "detected_only_weighted_mass_ug_L": low,
                    "official_group_weighted_mass_ug_L": official_total,
                    "detected_events": int(events[f"{s}_detected"].sum()),
                    "systems_with_detection": int(events.loc[events[f"{s}_detected"], "PWSID"].nunique()),
                    "bootstrap_replicates": BOOTSTRAPS,
                    "bootstrap_cluster": "PWSID",
                })
    return pd.DataFrame(rows)


def distribution_summary(events: pd.DataFrame, systems: pd.DataFrame) -> pd.DataFrame:
    rows = []
    estimators = {
        "official_group_totals": "official",
        "species_nd0": "lower",
        "species_nd_half_mrl": "mid",
        "censor_interval_lower": "interval_lower",
        "censor_interval_upper": "interval_upper",
    }
    for scope, frame in (("event", events), ("system_mean", systems)):
        for target in ("Extra4", "HAA6Br"):
            for estimator, suffix in estimators.items():
                col = f"{target}_fraction_{suffix}"
                values = frame[col].replace([np.inf, -np.inf], np.nan).dropna().to_numpy(dtype=float)
                if scope == "event":
                    finite = frame[col].replace([np.inf, -np.inf], np.nan).notna()
                    represented_systems = int(frame.loc[finite, "PWSID"].nunique())
                else:
                    represented_systems = len(values)
                rows.append({
                    "target_fraction": f"{target}/HAA9",
                    "scope": scope,
                    "estimator": estimator,
                    "n": len(values),
                    "represented_systems": represented_systems,
                    "mean": float(np.mean(values)),
                    "sd": float(np.std(values, ddof=1)),
                    "p05": float(np.quantile(values, 0.05)),
                    "p25": float(np.quantile(values, 0.25)),
                    "median": float(np.quantile(values, 0.50)),
                    "p75": float(np.quantile(values, 0.75)),
                    "p95": float(np.quantile(values, 0.95)),
                })
    return pd.DataFrame(rows)


def primary_reconciliation(events: pd.DataFrame) -> pd.DataFrame:
    rows = []
    mappings = {
        "HAA5": "HAA5_species_lower_ug_L",
        "HAA6Br": "HAA6Br_species_lower_ug_L",
        "HAA9": "HAA9_species_lower_ug_L",
    }
    for group, derived_col in mappings.items():
        official_col = f"{group}_official_ug_L"
        diff = events[derived_col] - events[official_col]
        matched_1e9 = diff.abs() <= 1e-9
        rows.append({
            "group": group,
            "events": len(events),
            "literal_binary_float_equal_matches": int((diff == 0).sum()),
            "matches_within_1e_9_ug_L": int(matched_1e9.sum()),
            "within_1e_6_ug_L": int((diff.abs() <= 1e-6).sum()),
            "mean_absolute_difference_ug_L": float(diff.abs().mean()),
            "maximum_absolute_difference_ug_L": float(diff.abs().max()),
            "nonmatching_events_gt_1e_9_ug_L": int((~matched_1e9).sum()),
        })
    return pd.DataFrame(rows)


def figure_source(events: pd.DataFrame, systems: pd.DataFrame, contributions: pd.DataFrame, distributions: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for _, r in contributions.iterrows():
        rows.append({
            "panel": "a" if r["target_group"] == "Extra4" else "b",
            "record_type": "species_contribution",
            "target": r["target_group"],
            "scope": r["weighting"],
            "label": r["abbreviation"],
            "value": r["primary_detect_only_share_official_denominator"],
            "ci_low": r["primary_cluster_bootstrap_ci_low"],
            "ci_high": r["primary_cluster_bootstrap_ci_high"],
            "censor_low": r["censor_interval_share_lower"],
            "censor_high": r["censor_interval_share_upper"],
            "n": r["events"],
            "systems": r["systems"],
        })
    for scope, frame in (("event", events), ("system_mean", systems)):
        for target in ("Extra4", "HAA6Br"):
            col = f"{target}_fraction_official"
            values = frame[col].dropna()
            represented_systems = int(frame.loc[values.index, "PWSID"].nunique()) if scope == "event" else len(values)
            for idx, value in values.items():
                rows.append({
                    "panel": "c",
                    "record_type": "official_fraction",
                    "target": target,
                    "scope": scope,
                    "label": str(idx),
                    "value": value,
                    "ci_low": math.nan,
                    "ci_high": math.nan,
                    "censor_low": math.nan,
                    "censor_high": math.nan,
                    "n": len(values),
                    "systems": represented_systems,
                })
    for target in ("Extra4", "HAA6Br"):
        for scope in ("event", "system_mean"):
            sub = distributions[(distributions["target_fraction"] == f"{target}/HAA9") & (distributions["scope"] == scope)]
            med = dict(zip(sub["estimator"], sub["median"]))
            nrow = sub[sub["estimator"] == "official_group_totals"].iloc[0]
            rows.append({
                "panel": "d",
                "record_type": "censor_sensitivity_median",
                "target": target,
                "scope": scope,
                "label": f"{target}_{scope}",
                "value": med["official_group_totals"],
                "ci_low": math.nan,
                "ci_high": math.nan,
                "censor_low": med["censor_interval_lower"],
                "censor_high": med["censor_interval_upper"],
                "n": int(nrow["n"]),
                "systems": int(nrow["represented_systems"]),
            })
    return pd.DataFrame(rows)


def configure_plotting() -> None:
    mpl.rcParams.update({
        "font.family": "Arial Narrow",
        "font.sans-serif": ["Arial Narrow"],
        "font.size": 8.2,
        "axes.titlesize": 9.0,
        "axes.labelsize": 8.5,
        "xtick.labelsize": 7.6,
        "ytick.labelsize": 7.6,
        "legend.fontsize": 7.2,
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
        "svg.fonttype": "none",
        "axes.linewidth": 0.7,
        "xtick.major.width": 0.6,
        "ytick.major.width": 0.6,
        "savefig.facecolor": "white",
    })


def draw_contribution_panel(ax, contributions: pd.DataFrame, target: str, species_order: list[str]) -> None:
    sub = contributions[contributions["target_group"] == target]
    y = np.arange(len(species_order))[::-1]
    offsets = {"event_mass_weighted": 0.13, "equal_system_mass_weighted": -0.13}
    markers = {"event_mass_weighted": "o", "equal_system_mass_weighted": "s"}
    for weighting in offsets:
        for i, abbr in enumerate(species_order):
            r = sub[(sub["weighting"] == weighting) & (sub["abbreviation"] == abbr)].iloc[0]
            yy = y[i] + offsets[weighting]
            point = 100 * r["primary_detect_only_share_official_denominator"]
            blo, bhi = 100 * r["primary_cluster_bootstrap_ci_low"], 100 * r["primary_cluster_bootstrap_ci_high"]
            clo, chi = 100 * r["censor_interval_share_lower"], 100 * r["censor_interval_share_upper"]
            ax.plot([clo, chi], [yy, yy], color=COLORS[abbr], lw=4.5, alpha=0.25, solid_capstyle="round", zorder=1)
            ax.plot([blo, bhi], [yy, yy], color="#222222", lw=0.8, zorder=2)
            ax.scatter(point, yy, s=28, marker=markers[weighting], facecolor=COLORS[abbr], edgecolor="#222222", linewidth=0.45, zorder=3)
    ax.set_yticks(y, species_order)
    ax.set_xlim(left=0)
    xmax = max(100 * sub["censor_interval_share_upper"].max() * 1.08, 10)
    ax.set_xlim(0, min(100, math.ceil(xmax / 5) * 5))
    ax.set_xlabel("Contribution to group mass (%)")
    ax.grid(axis="x", color="#D8D8D8", lw=0.5)
    ax.spines[["top", "right"]].set_visible(False)


def make_figure(events: pd.DataFrame, systems: pd.DataFrame, contributions: pd.DataFrame, distributions: pd.DataFrame) -> list[Path]:
    configure_plotting()
    fig, axes = plt.subplots(2, 2, figsize=(7.2, 7.25), constrained_layout=True)
    axa, axb, axc, axd = axes.ravel()

    draw_contribution_panel(axa, contributions, "Extra4", ["BCAA", "BDCAA", "CDBAA", "TBAA"])
    axa.set_title("Extra4 species share", loc="left", fontweight="bold")
    legend = [
        Line2D([0], [0], marker="o", color="none", markerfacecolor="#777777", markeredgecolor="#222222", markersize=5, label="Event weighted"),
        Line2D([0], [0], marker="s", color="none", markerfacecolor="#777777", markeredgecolor="#222222", markersize=5, label="Equal PWS weighted"),
        Line2D([0], [0], color="#777777", alpha=0.35, lw=4.5, label="Censor interval"),
    ]
    axa.legend(handles=legend, frameon=False, loc="lower right")

    draw_contribution_panel(axb, contributions, "HAA6Br", ["MBAA", "DBAA", "BCAA", "BDCAA", "CDBAA", "TBAA"])
    axb.set_title("HAA6Br species share", loc="left", fontweight="bold")

    arrays = [
        events["Extra4_fraction_official"].dropna().to_numpy(),
        systems["Extra4_fraction_official"].dropna().to_numpy(),
        events["HAA6Br_fraction_official"].dropna().to_numpy(),
        systems["HAA6Br_fraction_official"].dropna().to_numpy(),
    ]
    positions = np.arange(1, 5)
    vp = axc.violinplot(arrays, positions=positions, widths=0.78, showextrema=False)
    fills = ["#8CB8D8", "#4F84AA", "#E7B26A", "#CB7C18"]
    for body, color in zip(vp["bodies"], fills):
        body.set_facecolor(color)
        body.set_edgecolor("#333333")
        body.set_linewidth(0.55)
        body.set_alpha(0.9)
    for pos, values in zip(positions, arrays):
        q25, med, q75 = np.quantile(values, [0.25, 0.5, 0.75])
        axc.plot([pos, pos], [q25, q75], color="white", lw=3.0, solid_capstyle="round", zorder=3)
        axc.scatter(pos, med, s=16, facecolor="#222222", edgecolor="white", linewidth=0.4, zorder=4)
    axc.set_xticks(positions, ["Extra4\nevents", "Extra4\nPWS", "HAA6Br\nevents", "HAA6Br\nPWS"])
    axc.set_ylim(0, 1)
    axc.yaxis.set_major_formatter(mpl.ticker.PercentFormatter(1.0))
    axc.set_ylabel("Fraction of official HAA9")
    axc.set_title("Official composition fractions", loc="left", fontweight="bold")
    axc.grid(axis="y", color="#D8D8D8", lw=0.5)
    axc.spines[["top", "right"]].set_visible(False)
    for pos, values in zip(positions, arrays):
        axc.text(
            pos, 0.975, f"n={len(values):,}", ha="center", va="top", fontsize=6.9,
            bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.78, "pad": 0.7},
            clip_on=True, zorder=6,
        )

    labels = [
        ("Extra4", "event", "Extra4, events", "#4F84AA"),
        ("Extra4", "system_mean", "Extra4, PWS", "#4F84AA"),
        ("HAA6Br", "event", "HAA6Br, events", "#CB7C18"),
        ("HAA6Br", "system_mean", "HAA6Br, PWS", "#CB7C18"),
    ]
    y = np.arange(4)[::-1]
    for yy, (target, scope, label, color) in zip(y, labels):
        sub = distributions[(distributions["target_fraction"] == f"{target}/HAA9") & (distributions["scope"] == scope)]
        med = dict(zip(sub["estimator"], sub["median"]))
        lo, hi = med["censor_interval_lower"], med["censor_interval_upper"]
        midpoint = med["species_nd_half_mrl"]
        official = med["official_group_totals"]
        axd.plot([lo, hi], [yy, yy], color=color, alpha=0.35, lw=7, solid_capstyle="round")
        axd.scatter(midpoint, yy, s=34, marker="o", facecolor=color, edgecolor="#222222", linewidth=0.5, label=None, zorder=3)
        axd.scatter(official, yy, s=38, marker="D", facecolor="white", edgecolor="#111111", linewidth=0.75, zorder=4)
    axd.set_yticks(y, [x[2] for x in labels])
    axd.set_xlim(0, 1)
    axd.xaxis.set_major_formatter(mpl.ticker.PercentFormatter(1.0))
    axd.set_xlabel("Median fraction of HAA9")
    axd.set_title("Censoring sensitivity", loc="left", fontweight="bold")
    axd.grid(axis="x", color="#D8D8D8", lw=0.5)
    axd.spines[["top", "right"]].set_visible(False)
    axd.legend(handles=[
        Line2D([0], [0], marker="D", color="none", markerfacecolor="white", markeredgecolor="#111111", markersize=5, label="Official totals"),
        Line2D([0], [0], marker="o", color="none", markerfacecolor="#777777", markeredgecolor="#222222", markersize=5, label="ND = 1/2 MRL"),
        Line2D([0], [0], color="#777777", alpha=0.35, lw=6, label="Censor interval"),
    ], frameon=False, loc="lower right")

    for label, ax in zip(("(a)", "(b)", "(c)", "(d)"), (axa, axb, axc, axd)):
        ax.text(-0.13, 1.07, label, transform=ax.transAxes, fontsize=10, fontweight="bold", va="top", ha="left")

    fig.suptitle("Species-resolved HAA composition in UCMR4", fontsize=11, fontweight="bold")
    paths = [
        FIG / "FigS6_species_level_HAA_composition.svg",
        FIG / "FigS6_species_level_HAA_composition.pdf",
        FIG / "FigS6_species_level_HAA_composition.png",
    ]
    fig.savefig(paths[0], bbox_inches="tight")
    fig.savefig(paths[1], bbox_inches="tight")
    fig.savefig(paths[2], dpi=600, bbox_inches="tight")
    plt.close(fig)
    return paths


def methods_results_note(audit: dict, coverage: pd.DataFrame, contributions: pd.DataFrame, distributions: pd.DataFrame, reconciliation: pd.DataFrame) -> str:
    def dist(target, scope, estimator="official_group_totals"):
        return distributions[(distributions.target_fraction == target) & (distributions.scope == scope) & (distributions.estimator == estimator)].iloc[0]

    def share(target, abbr, weighting):
        return contributions[(contributions.target_group == target) & (contributions.abbreviation == abbr) & (contributions.weighting == weighting)].iloc[0]

    extra_event = dist("Extra4/HAA9", "event")
    extra_system = dist("Extra4/HAA9", "system_mean")
    br_event = dist("HAA6Br/HAA9", "event")
    br_system = dist("HAA6Br/HAA9", "system_mean")
    primary_extra = contributions[(contributions.target_group == "Extra4") & (contributions.weighting == "event_mass_weighted")].sort_values("primary_detect_only_share_official_denominator", ascending=False)
    primary_br = contributions[(contributions.target_group == "HAA6Br") & (contributions.weighting == "event_mass_weighted")].sort_values("primary_detect_only_share_official_denominator", ascending=False)
    extra_text = ", ".join(f"{r.abbreviation} {100*r.primary_detect_only_share_official_denominator:.1f}%" for r in primary_extra.itertuples())
    br_text = ", ".join(f"{r.abbreviation} {100*r.primary_detect_only_share_official_denominator:.1f}%" for r in primary_br.itertuples())
    rec_text = "; ".join(f"{r.group}: {int(r.matches_within_1e_9_ug_L):,}/{int(r.events):,} within 10^-9 µg/L" for r in reconciliation.itertuples())
    return f"""# Species-level UCMR4 addendum

## Methods

Species-level UCMR4 records were read from the Peterson et al. repository frozen at commit `{COMMIT}`. The four FOIA workbooks contained {audit['species']['rows']:,} records from {audit['species']['systems']:,} public water systems. Events were linked to EPA's official UCMR4 HAA5, HAA6Br, and HAA9 rows using PWSID, facility, sampling point, sample ID, collection date, and sample-event code. Repository codes `SEH1` to `SEH4` were normalized to the official `SE1` to `SE4` convention. Analysis required all nine species and all three official group totals at the same event.

Official EPA group totals were primary. Species results marked `=` were represented by their reported concentration. Results marked `<` were interval-censored from 0 to the species MRL. The primary attribution assigned nondetects to zero, matching the public repository notebook. Sensitivity used MRL/2 substitution and species-specific interval bounds. For a species share, the lower bound divided that species' lower mass by its lower mass plus the upper mass of all other group members; the upper bound used its upper mass plus the lower mass of the others. Event-mass summaries weighted each complete event equally. Equal-system summaries first averaged concentrations within PWS and then weighted each PWS equally. Statistical intervals for primary shares used {BOOTSTRAPS:,} PWS-cluster bootstrap replicates with seed {SEED}.

Extra4 was defined as official HAA9 minus official HAA5 and attributed to BCAA, BDCAA, CDBAA, and TBAA. Brominated composition used official HAA6Br and the six brominated species MBAA, DBAA, BCAA, BDCAA, CDBAA, and TBAA. Group fractions were calculated only where official HAA9 was positive. Primary system fractions averaged eligible event fractions within PWS, matching the main UCMR4 analysis. Ratios of within-system mean group concentrations were retained as separate sensitivity fields.

## Results

The linked cohort contained {audit['cohort']['analysis_events']:,} complete nine-species events from {audit['cohort']['analysis_systems']:,} systems spanning {audit['cohort']['date_min']} to {audit['cohort']['date_max']}. Every complete event linked to all three official group totals. Literal equality is inappropriate for decimal sums stored as binary floating-point values and flags thousands of numerically negligible differences. With a 10^-9 µg/L tolerance, detect-only species sums reproduced the official results almost exactly ({rec_text}). Only six of 190,296 group comparisons exceeded 10^-6 µg/L. Those six comparisons came from three physical samples, and each affected sample differed by exactly one relevant species MRL. Official group values were retained without alteration.

Among {int(extra_event['n']):,} events with positive official HAA9, the median Extra4/HAA9 fraction was {extra_event['median']:.3f} (IQR {extra_event['p25']:.3f} to {extra_event['p75']:.3f}). The corresponding within-system mean event fraction was {extra_system['median']:.3f} (IQR {extra_system['p25']:.3f} to {extra_system['p75']:.3f}; {int(extra_system['n']):,} systems). Event-mass contributions to official Extra4 were {extra_text}. The species-specific censor intervals are reported in `species_group_contributions.csv`; they quantify unresolved sub-MRL mass rather than treating nondetects as observed zero.

The median official HAA6Br/HAA9 fraction was {br_event['median']:.3f} at the event level and {br_system['median']:.3f} after within-system averaging. Event-mass contributions to official HAA6Br were {br_text}. Results were reported separately under event and equal-system weighting because frequent samplers otherwise contribute more mass to national aggregates. These fractions describe the composition of measured HAAs, not toxicological contribution, health-risk coverage, exposure, or regulatory compliance.

## Figure caption

**Supplementary Fig. 6 | Species-resolved composition of UCMR4 haloacetic acids.** (a) Contributions of BCAA, BDCAA, CDBAA, and TBAA to Extra4, defined from official HAA9 minus HAA5. (b) Contributions of six brominated species to official HAA6Br. Points show detect-only attribution using official group totals; thin lines are 95% PWS-cluster bootstrap intervals; broad bands are species-specific censor intervals. Event weighting and equal-PWS weighting are shown separately. (c) Distributions of official Extra4/HAA9 and HAA6Br/HAA9 fractions for positive-HAA9 events and within-system mean event fractions. White bars denote interquartile ranges and dark points denote medians. (d) Official medians compared with MRL/2 substitution and interval-censor bounds. Fractions characterize measured mass and do not represent toxicity, exposure, or violations.
"""


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    FIG.mkdir(parents=True, exist_ok=True)

    species_events, species_audit = load_species_events()
    official_groups, official_audit = load_official_groups()
    events, cohort_audit = build_event_frame(species_events, official_groups)
    systems = build_system_frame(events)
    coverage = species_coverage(events)
    contributions = contribution_table(events, systems)
    distributions = distribution_summary(events, systems)
    reconciliation = primary_reconciliation(events)

    outputs = []
    outputs.extend(save_table(events, "ucmr4_species_event_panel"))
    outputs.extend(save_table(systems, "ucmr4_species_system_summary"))
    outputs.extend(save_table(coverage, "species_coverage_censoring"))
    outputs.extend(save_table(contributions, "species_group_contributions"))
    outputs.extend(save_table(distributions, "species_composition_distributions"))
    outputs.extend(save_table(reconciliation, "species_official_group_reconciliation"))

    fig_source = figure_source(events, systems, contributions, distributions)
    outputs.extend(save_table(fig_source, "FigS6_source_data"))
    figure_paths = make_figure(events, systems, contributions, distributions)
    outputs.extend(figure_paths)

    audit = {
        "generated_utc": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "analysis": "species-level UCMR4 HAA composition addendum",
        "species_repository": {
            "url": "https://github.com/wraseman/haa-cytotoxicity-occurrence",
            "commit": COMMIT,
            "license": "MIT",
            "associated_doi": "10.1002/aws2.1351",
        },
        "official_ucmr4_source": {
            "url": "https://www.epa.gov/system/files/other-files/2024-04/ucmr4-occurrence-data.zip",
            "path": str(UCMR4),
            "sha256": sha256(UCMR4),
        },
        "species": species_audit,
        "official": official_audit,
        "cohort": cohort_audit,
        "censoring": {
            "primary": "detected value; nondetect = 0 for attribution only",
            "midpoint": "nondetect = MRL/2",
            "interval": "nondetect in [0, MRL] with species-specific ratio bounds",
        },
        "weighting": {
            "event": "each complete event equally weighted",
            "system": "within-PWS means, then each PWS equally weighted",
        },
        "bootstrap": {"replicates": BOOTSTRAPS, "seed": SEED, "cluster": "PWSID"},
        "headline": {
            "official_composition": distributions[
                distributions.estimator == "official_group_totals"
            ].to_dict(orient="records"),
            "event_mass_contributions": contributions[
                contributions.weighting == "event_mass_weighted"
            ].to_dict(orient="records"),
            "reconciliation": reconciliation.to_dict(orient="records"),
        },
    }
    note = methods_results_note(audit, coverage, contributions, distributions, reconciliation)
    note_path = OUT / "METHODS_RESULTS_SPECIES_ADDENDUM.md"
    note_path.write_text(note, encoding="utf-8")
    outputs.append(note_path)

    manifest_rows = []
    for path in outputs:
        manifest_rows.append({
            "path": str(path),
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
        })
    manifest = pd.DataFrame(manifest_rows)
    manifest_path = OUT / "species_addendum_manifest.csv"
    manifest.to_csv(manifest_path, index=False)
    audit["outputs"] = manifest_rows
    summary_path = OUT / "species_analysis_summary.json"
    summary_path.write_text(json.dumps(audit, indent=2, ensure_ascii=False), encoding="utf-8")

    print(json.dumps({
        "events": len(events),
        "systems": len(systems),
        "results_directory": str(OUT),
        "figure": [str(p) for p in figure_paths],
        "summary": str(summary_path),
    }, indent=2))


if __name__ == "__main__":
    main()
